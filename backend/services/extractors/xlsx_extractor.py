"""Excel text extraction using openpyxl."""
from openpyxl import load_workbook


def extract_xlsx(file_path: str) -> str:
    """Extract text from an Excel file."""
    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        text_parts = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_text = [f"[Sheet: {sheet_name}]"]

            for row in ws.iter_rows(max_row=500, values_only=True):
                row_vals = [str(c) if c is not None else "" for c in row]
                row_str = " | ".join(v for v in row_vals if v)
                if row_str.strip():
                    sheet_text.append(row_str)

            if len(sheet_text) > 1:
                text_parts.append("\n".join(sheet_text))

        wb.close()
        return "\n\n".join(text_parts)
    except Exception as e:
        raise Exception(f"XLSX extraction failed: {e}")
