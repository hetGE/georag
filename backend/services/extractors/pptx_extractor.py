"""PowerPoint text extraction using python-pptx."""
from pptx import Presentation


def extract_pptx(file_path: str) -> str:
    """Extract text from a .pptx file."""
    try:
        prs = Presentation(file_path)
        text_parts = []

        for i, slide in enumerate(prs.slides, 1):
            slide_text = [f"[Slide {i}]"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_text.append(text)
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(
                            cell.text.strip() for cell in row.cells if cell.text.strip()
                        )
                        if row_text:
                            slide_text.append(row_text)

            if len(slide_text) > 1:
                text_parts.append("\n".join(slide_text))

        return "\n\n".join(text_parts)
    except Exception as e:
        raise Exception(f"PPTX extraction failed: {e}")
