"""Word document text extraction using python-docx."""
from docx import Document


async def extract_docx(file_path: str) -> str:
    """Extract text from a .docx file."""
    try:
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        return "\n\n".join(paragraphs)
    except Exception as e:
        raise Exception(f"DOCX extraction failed: {e}")
